from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
STEM = "Introduction_Methodes_Resultats_preliminaires_Comparaison_UNI_Virchow_2026-10-17"
PDF_PATH = ROOT / "output" / "pdf" / f"{STEM}.pdf"
DOCX_PATH = ROOT / "output" / "docs" / f"{STEM}.docx"


pdf = PdfReader(PDF_PATH)
pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
pdf_checks = {
    "pages_expected_range": 10 <= len(pdf.pages) <= 12,
    "objective": "objectif principal est de comparer notre modèle supervisé à UNI et à Virchow" in pdf_text,
    "preparatory_scope": "ne répondent pas encore" in pdf_text and "objectif principal" in pdf_text,
    "uni_virchow": "UNI et Virchow" in pdf_text,
    "iter8_metrics": "0,579" in pdf_text and "0,966" in pdf_text,
    "junior_review": "28 contenaient de" in pdf_text and "tumeur au centre" in pdf_text,
    "frozen_rule": "seuil patch de 0,25" in pdf_text and "sensibilit" in pdf_text and "1,000" in pdf_text,
    "references": "doi:10.1038/s41591-024-02857-3" in pdf_text,
    "no_raw_citation_keys": "@chen" not in pdf_text and "@who" not in pdf_text,
}

docx = Document(DOCX_PATH)
docx_text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
docx_checks = {
    "title_style": docx.paragraphs[0].style.name == "Title",
    "tables_expected": len(docx.tables) == 5,
    "objective": "objectif principal est de comparer notre modèle supervisé à UNI et à Virchow" in docx_text,
    "preparatory_scope": "ne répondent pas encore à l'objectif principal" in docx_text,
    "junior_review": "28 contenaient de" in docx_text and "tumeur au centre" in docx_text,
    "frozen_rule": "seuil patch de 0,25" in docx_text and "sensibilité de 1,000" in docx_text,
    "no_raw_citation_keys": "@chen" not in docx_text and "@who" not in docx_text,
    "no_duplicate_title": docx_text.count(
        "Comparaison d'un algorithme supervisé et des modèles fondationnels UNI et Virchow"
    )
    == 1,
}

print({"pdf": pdf_checks, "docx": docx_checks})
if not all(pdf_checks.values()) or not all(docx_checks.values()):
    raise SystemExit(1)
