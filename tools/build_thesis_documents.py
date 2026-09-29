from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "these_intro_methodes_brouillon.md"
DOCX_OUT = ROOT / "output" / "docs" / "Introduction_Methodes_Resultats_preliminaires_Comparaison_UNI_Virchow_2026-09-29.docx"
PDF_OUT = ROOT / "output" / "pdf" / "Introduction_Methodes_Resultats_preliminaires_Comparaison_UNI_Virchow_2026-09-29.pdf"

TITLE = (
    "Comparaison d'un algorithme supervisé et des modèles fondationnels UNI et Virchow "
    "pour la détection de métastases ganglionnaires mammaires sur les lames "
    "histologiques numérisées du service"
)
SUBTITLE = "Introduction, méthodes et résultats préliminaires"
VERSION = "Version de travail — 29 septembre 2026"

NAVY = "1F4E78"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9E1E8"
DARK_GRAY = "404040"

CITATIONS = {
    "whoBreastCancer2026": 1,
    "incaCancerSein2026": 2,
    "spfCancerSein2026": 3,
    "aiLymphNodeMetastasesReview2023": 4,
    "rapportNumerisationACP2025": 5,
    "solimanBreastAI2024": 6,
    "katayamaBreastPathologyAI2024": 7,
    "ehteshamiBejnordiCamelyon2017": 8,
    "digitalPathologyRoutineReview2024": 9,
    "chenUNI2024": 10,
    "vorontsovVirchow2024": 11,
}

REFERENCES = [
    "Organisation mondiale de la Santé. Breast cancer. Mise à jour du 3 juillet 2026. https://www.who.int/news-room/fact-sheets/detail/breast-cancer",
    "Institut national du cancer. Les cancers du sein. Mise à jour du 22 juillet 2026. https://en-www.cancer.fr/professionnels-de-sante/statistiques-et-chiffres-sur-les-cancers/epidemiologie-des-cancers/cancer-du-sein",
    "Santé publique France. Cancer du sein — Données. Mise à jour du 6 juillet 2026. https://www.santepubliquefrance.fr/index.php/cancer-du-sein/donnees",
    "Value of Artificial Intelligence in Evaluating Lymph Node Metastases: A Systematic Review. 2023. https://pmc.ncbi.nlm.nih.gov/articles/PMC10177013/",
    "Ministère chargé de la Santé. Rapport de la mission ministérielle sur la politique de numérisation de l'anatomie et cytologie pathologiques. 2025.",
    "Soliman A, Li Z, Parwani AV. Artificial intelligence's impact on breast cancer pathology. Diagnostic Pathology. 2024. doi:10.1186/s13000-024-01453-w.",
    "Katayama A, Aoki Y, Watanabe Y, et al. Current status and prospects of artificial intelligence in breast cancer pathology. International Journal of Clinical Oncology. 2024;29(11):1648–1668. doi:10.1007/s10147-024-02513-3.",
    "Ehteshami Bejnordi B, Veta M, van Diest PJ, et al. Diagnostic Assessment of Deep Learning Algorithms for Detection of Lymph Node Metastases in Women With Breast Cancer. JAMA. 2017;318(22):2199–2210. doi:10.1001/jama.2017.14585.",
    "Implementation of Digital Pathology and Artificial Intelligence in Routine Pathology Practice. Laboratory Investigation. 2024. doi:10.1016/j.labinv.2024.102111.",
    "Chen RJ, Ding T, Lu MY, Williamson DFK, et al. Towards a general-purpose foundation model for computational pathology. Nature Medicine. 2024;30:850–862. doi:10.1038/s41591-024-02857-3.",
    "Vorontsov E, et al. A foundation model for clinical-grade computational pathology and rare cancers detection. Nature Medicine. 2024. doi:10.1038/s41591-024-03141-0.",
]


@dataclass
class Block:
    kind: str
    text: str | None = None
    level: int | None = None
    rows: list[list[str]] | None = None


def replace_citations(text: str) -> str:
    def repl(match: re.Match[str]) -> str:
        keys = re.findall(r"@([A-Za-z0-9_]+)", match.group(1))
        nums = [CITATIONS[k] for k in keys if k in CITATIONS]
        return "[" + ", ".join(str(n) for n in nums) + "]" if nums else match.group(0)

    return re.sub(r"\[([^\]]*@[A-Za-z0-9_]+[^\]]*)\]", repl, text)


def clean_inline(text: str) -> str:
    return replace_citations(text.strip())


def parse_markdown() -> list[Block]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    para: list[str] = []
    in_code = False

    def flush_para() -> None:
        nonlocal para
        if para:
            blocks.append(Block("paragraph", clean_inline(" ".join(s.strip() for s in para))))
            para = []

    i = 0
    skip_bibliography = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("## Bibliographie de travail"):
            flush_para()
            skip_bibliography = True
            break
        if i == 0 and stripped.startswith("# "):
            i += 1
            continue
        if stripped.startswith("*Version de travail"):
            i += 1
            continue
        if stripped == "## Titre de travail":
            flush_para()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("## "):
                i += 1
            continue
        if stripped.startswith("```"):
            flush_para()
            in_code = not in_code
            i += 1
            continue
        if in_code:
            if stripped:
                blocks.append(Block("code", stripped))
            i += 1
            continue
        if not stripped:
            flush_para()
            i += 1
            continue
        if stripped.startswith("#"):
            flush_para()
            m = re.match(r"^(#{2,4})\s+(.+)$", stripped)
            if m:
                blocks.append(Block("heading", clean_inline(m.group(2)), len(m.group(1)) - 1))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip().lstrip("|")):
            flush_para()
            rows: list[list[str]] = []
            rows.append([clean_inline(c.strip()) for c in stripped.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([clean_inline(c.strip()) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if stripped.startswith("- "):
            flush_para()
            blocks.append(Block("bullet", clean_inline(stripped[2:])))
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not lines[i].strip().startswith("-"):
                blocks[-1].text += " " + clean_inline(lines[i])
                i += 1
            continue
        para.append(stripped)
        i += 1
    flush_para()
    blocks.append(Block("heading", "Références", 1))
    for idx, ref in enumerate(REFERENCES, 1):
        blocks.append(Block("reference", f"{idx}. {ref}"))
    return blocks


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_inline_runs(paragraph, text: str, *, size: float | None = None) -> None:
    parts = re.split(r"(`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((size or 10.8) - 0.5)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            if size:
                run.font.size = Pt(size)
        else:
            run = paragraph.add_run(part)
            if size:
                run.font.size = Pt(size)


def configure_docx_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("202020")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)

    for name, size, before, after in (
        ("Heading 1", 15, 12, 8),
        ("Heading 2", 12.5, 10, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Reference" not in styles:
        ref = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref = styles["Reference"]
    ref.font.name = "Times New Roman"
    ref.font.size = Pt(9.5)
    ref.paragraph_format.left_indent = Cm(0.65)
    ref.paragraph_format.first_line_indent = Cm(-0.65)
    ref.paragraph_format.space_after = Pt(4)
    ref.paragraph_format.line_spacing = 1.05


def configure_section(section, *, first: bool = False) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = first


def add_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = "THÈSE — PATHOLOGIE NUMÉRIQUE"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Page ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    add_field(p, "PAGE")


def add_docx_table(doc: Document, rows: list[list[str]]) -> None:
    chunks = [rows]
    if len(rows[0]) == 10 and rows[0][0] == "Modèle":
        chunks = [
            [[r[i] for i in range(6)] for r in rows],
            [[r[0]] + [r[i] for i in range(6, 10)] for r in rows],
        ]
    for ci, data in enumerate(chunks):
        if ci:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("Matrice de confusion")
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
        table = doc.add_table(rows=len(data), cols=len(data[0]))
        table.alignment = 1
        table.autofit = True
        set_repeat_table_header(table.rows[0])
        for r_idx, row in enumerate(data):
            for c_idx, value in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_border(cell)
                if r_idx == 0:
                    set_cell_shading(cell, NAVY)
                elif r_idx % 2 == 1:
                    set_cell_shading(cell, PALE_BLUE)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                add_inline_runs(p, value, size=8.4 if len(data[0]) > 5 else 9)
                for run in p.runs:
                    run.font.name = "Arial"
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx(blocks: list[Block]) -> None:
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_docx_styles(doc)
    configure_section(doc.sections[0], first=True)
    add_header_footer(doc.sections[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(22)
    p.style = doc.styles["Title"]
    p.add_run(TITLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(SUBTITLE)
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(VERSION)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string("555555")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Document de travail destiné à la relecture senior")
    r.italic = True
    r.font.size = Pt(10.5)
    doc.add_page_break()

    p = doc.add_paragraph("Points proposés pour la relecture", style="Heading 1")
    prompts = [
        "Le cadrage scientifique et la hiérarchie des objectifs sont-ils adaptés : comparaison principale sur les lames du service, avec CAMELYON16 comme phase préparatoire ?",
        "Le compromis observé entre sensibilité et spécificité pour iter7 et iter8 justifie-t-il l'expérience d'enrichissement équilibré proposée avant de figer le comparateur supervisé ?",
        "Le protocole envisagé permet-il une comparaison suffisamment équitable entre le modèle supervisé, UNI et Virchow sur la cohorte locale ?",
    ]
    intro = doc.add_paragraph()
    intro.add_run("Cette version est volontairement structurée pour discuter en priorité de l'objectif principal et du passage à l'étude comparative locale.")
    for item in prompts:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        add_inline_runs(p, item)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Statut du travail présenté")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        "Les résultats chiffrés portent actuellement sur la phase CAMELYON16. La comparaison avec UNI et Virchow sur les lames du service constitue l'étude principale à venir.",
    )

    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                if block.text == "Fil conducteur":
                    doc.add_page_break()
                style = "Heading 1"
            elif block.level == 2:
                style = "Heading 2"
            else:
                style = "Heading 3"
            doc.add_paragraph(block.text, style=style)
        elif block.kind == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.55)
            if (block.text or "").startswith("Le critère principal sera défini"):
                p.paragraph_format.keep_together = True
            add_inline_runs(p, block.text or "")
        elif block.kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline_runs(p, block.text or "")
        elif block.kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(block.text or "")
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif block.kind == "table":
            add_docx_table(doc, block.rows or [])
        elif block.kind == "reference":
            p = doc.add_paragraph(style="Reference")
            add_inline_runs(p, block.text or "", size=9.5)

    props = doc.core_properties
    props.title = TITLE
    props.subject = SUBTITLE
    props.keywords = "pathologie numérique, CAMELYON16, UNI, Virchow, métastases ganglionnaires"
    doc.save(DOCX_OUT)


def register_fonts() -> None:
    fonts = Path("C:/Windows/Fonts")
    pdfmetrics.registerFont(TTFont("TimesNewRoman", fonts / "times.ttf"))
    pdfmetrics.registerFont(TTFont("TimesNewRoman-Bold", fonts / "timesbd.ttf"))
    pdfmetrics.registerFont(TTFont("TimesNewRoman-Italic", fonts / "timesi.ttf"))
    pdfmetrics.registerFont(TTFont("Arial", fonts / "arial.ttf"))
    pdfmetrics.registerFont(TTFont("Arial-Bold", fonts / "arialbd.ttf"))
    pdfmetrics.registerFontFamily(
        "TimesNewRoman",
        normal="TimesNewRoman",
        bold="TimesNewRoman-Bold",
        italic="TimesNewRoman-Italic",
        boldItalic="TimesNewRoman-Bold",
    )


def pdf_markup(text: str) -> str:
    parts = re.split(r"(`[^`]+`|\*[^*]+\*)", text)
    rendered: list[str] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            rendered.append(f'<font name="Arial" size="8.5">{escape(part[1:-1])}</font>')
        elif part.startswith("*") and part.endswith("*"):
            rendered.append(f"<i>{escape(part[1:-1])}</i>")
        else:
            rendered.append(escape(part))
    return "".join(rendered)


class ThesisDocTemplate(BaseDocTemplate):
    def __init__(self, filename: str):
        super().__init__(
            filename,
            pagesize=A4,
            leftMargin=2.25 * cm,
            rightMargin=2.25 * cm,
            topMargin=2.1 * cm,
            bottomMargin=1.8 * cm,
            title=TITLE,
            author="",
            subject=SUBTITLE,
        )
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="body")
        self.addPageTemplates(
            [
                PageTemplate(id="cover", frames=[frame], onPage=self.draw_cover, autoNextPageTemplate="body"),
                PageTemplate(id="body", frames=[frame], onPage=self.draw_body),
            ]
        )

    @staticmethod
    def draw_cover(canvas, doc):
        canvas.saveState()
        canvas.setFont("Arial", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(A4[0] / 2, 0.9 * cm, VERSION)
        canvas.restoreState()

    @staticmethod
    def draw_body(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.setFont("Arial", 8)
        canvas.drawRightString(A4[0] - 2.25 * cm, A4[1] - 1.05 * cm, "THÈSE — PATHOLOGIE NUMÉRIQUE")
        canvas.drawCentredString(A4[0] / 2, 0.8 * cm, f"Page {doc.page}")
        canvas.restoreState()


def make_pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "TitleCustom",
            parent=styles["Title"],
            fontName="TimesNewRoman-Bold",
            fontSize=22,
            leading=27,
            textColor=colors.black,
            alignment=TA_CENTER,
            spaceAfter=18,
        ),
        "subtitle": ParagraphStyle(
            "SubtitleCustom",
            fontName="Arial-Bold",
            fontSize=14,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=14,
        ),
        "version": ParagraphStyle(
            "VersionCustom",
            fontName="Arial",
            fontSize=10.5,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#555555"),
        ),
        "h1": ParagraphStyle(
            "H1Custom",
            fontName="Arial-Bold",
            fontSize=15,
            leading=19,
            textColor=colors.black,
            spaceAfter=9,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "H2Custom",
            fontName="Arial-Bold",
            fontSize=12.5,
            leading=15,
            textColor=colors.black,
            spaceBefore=10,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "H3Custom",
            fontName="Arial-Bold",
            fontSize=11.5,
            leading=14,
            textColor=colors.black,
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "BodyCustom",
            fontName="TimesNewRoman",
            fontSize=10.6,
            leading=13.3,
            alignment=TA_JUSTIFY,
            firstLineIndent=0.55 * cm,
            spaceAfter=5,
            allowWidows=0,
            allowOrphans=0,
        ),
        "bullet": ParagraphStyle(
            "BulletCustom",
            fontName="TimesNewRoman",
            fontSize=10.6,
            leading=13.3,
            alignment=TA_JUSTIFY,
            leftIndent=0.8 * cm,
            firstLineIndent=-0.35 * cm,
            bulletIndent=0.25 * cm,
            spaceAfter=4,
        ),
        "code": ParagraphStyle(
            "CodeCustom",
            fontName="Arial",
            fontSize=9,
            leading=11,
            leftIndent=0.8 * cm,
            spaceAfter=2,
        ),
        "reference": ParagraphStyle(
            "ReferenceCustom",
            fontName="TimesNewRoman",
            fontSize=9.2,
            leading=11.2,
            leftIndent=0.65 * cm,
            firstLineIndent=-0.65 * cm,
            spaceAfter=4,
        ),
    }


def pdf_table(rows: list[list[str]], styles) -> list:
    chunks = [rows]
    if len(rows[0]) == 10 and rows[0][0] == "Modèle":
        chunks = [
            [[r[i] for i in range(6)] for r in rows],
            [[r[0]] + [r[i] for i in range(6, 10)] for r in rows],
        ]
    flow = []
    for ci, data in enumerate(chunks):
        if ci:
            flow.append(Spacer(1, 4))
            flow.append(Paragraph("Matrice de confusion", styles["h3"]))
        pstyle = ParagraphStyle("TableCell", fontName="Arial", fontSize=8, leading=9.5, alignment=TA_CENTER)
        hstyle = ParagraphStyle("TableHead", fontName="Arial-Bold", fontSize=7.7, leading=9, textColor=colors.white, alignment=TA_CENTER)
        pdata = [[Paragraph(pdf_markup(v), hstyle if ri == 0 else pstyle) for v in row] for ri, row in enumerate(data)]
        widths = [3.2 * cm] + [(15.9 * cm - 3.2 * cm) / (len(data[0]) - 1)] * (len(data[0]) - 1)
        table = Table(pdata, colWidths=widths, repeatRows=1, hAlign="CENTER")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#EAF1F8")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9E1E8")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                    ("ALIGN", (0, 1), (0, -1), "LEFT"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        flow.extend([table, Spacer(1, 5)])
    return flow


def build_pdf(blocks: list[Block]) -> None:
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    register_fonts()
    styles = make_pdf_styles()
    story = [
        Spacer(1, 3.2 * cm),
        Paragraph(pdf_markup(TITLE), styles["title"]),
        Spacer(1, 0.3 * cm),
        Paragraph(SUBTITLE, styles["subtitle"]),
        Paragraph(VERSION, styles["version"]),
        Spacer(1, 4.1 * cm),
        Paragraph("<i>Document de travail destiné à la relecture senior</i>", styles["version"]),
        PageBreak(),
        Paragraph("Points proposés pour la relecture", styles["h1"]),
        Paragraph(
            "Cette version est volontairement structurée pour discuter en priorité de l'objectif principal et du passage à l'étude comparative locale.",
            styles["body"],
        ),
    ]
    prompts = [
        "Le cadrage scientifique et la hiérarchie des objectifs sont-ils adaptés : comparaison principale sur les lames du service, avec CAMELYON16 comme phase préparatoire ?",
        "Le compromis observé entre sensibilité et spécificité pour iter7 et iter8 justifie-t-il l'expérience d'enrichissement équilibré proposée avant de figer le comparateur supervisé ?",
        "Le protocole envisagé permet-il une comparaison suffisamment équitable entre le modèle supervisé, UNI et Virchow sur la cohorte locale ?",
    ]
    for item in prompts:
        story.append(Paragraph("• " + pdf_markup(item), styles["bullet"]))
    story.extend(
        [
            Spacer(1, 8),
            Paragraph("Statut du travail présenté", styles["h2"]),
            Paragraph(
                "Les résultats chiffrés portent actuellement sur la phase CAMELYON16. La comparaison avec UNI et Virchow sur les lames du service constitue l'étude principale à venir.",
                styles["body"],
            ),
        ]
    )

    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                if block.text == "Fil conducteur":
                    story.append(PageBreak())
                style = styles["h1"]
            elif block.level == 2:
                style = styles["h2"]
            else:
                style = styles["h3"]
            story.append(Paragraph(pdf_markup(block.text or ""), style))
        elif block.kind == "paragraph":
            paragraph = Paragraph(pdf_markup(block.text or ""), styles["body"])
            if (block.text or "").startswith("Le critère principal sera défini"):
                story.append(KeepTogether([paragraph]))
            else:
                story.append(paragraph)
        elif block.kind == "bullet":
            story.append(Paragraph("• " + pdf_markup(block.text or ""), styles["bullet"]))
        elif block.kind == "code":
            story.append(Paragraph(pdf_markup(block.text or ""), styles["code"]))
        elif block.kind == "table":
            story.extend(pdf_table(block.rows or [], styles))
        elif block.kind == "reference":
            story.append(Paragraph(pdf_markup(block.text or ""), styles["reference"]))

    doc = ThesisDocTemplate(str(PDF_OUT))
    doc.build(story)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", action="store_true")
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    blocks = parse_markdown()
    if not args.docx and not args.pdf:
        args.docx = args.pdf = True
    if args.docx:
        build_docx(blocks)
        print(DOCX_OUT)
    if args.pdf:
        build_pdf(blocks)
        print(PDF_OUT)


if __name__ == "__main__":
    main()
